from __future__ import annotations

from collections.abc import Sequence
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
import math
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from .colors import ColorLike
from .coords import ColumnRef, cell_ref, column_to_index, index_to_column
from .styles import (
    make_alignment,
    make_border,
    make_fill,
    make_font,
    normalize_number_format,
    parse_text_style,
)

FillValue = str | int | float | bool | date | datetime | time | Decimal | None
RawFillValue = FillValue | object
SheetRef = int | str | Worksheet | None


@dataclass(frozen=True)
class ActiveSheet:
    """Small printable description of the current active worksheet."""

    index: int
    name: str

    def __str__(self) -> str:
        return f"{self.index}: {self.name}"

    def __iter__(self):
        yield self.index
        yield self.name




def nodia_casefold(s: str) -> str:
    """Casefold helper that also removes diacritics for forgiving header lookup."""
    s = (s or "").strip()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return s.casefold()


def _style_color_text(color: Any) -> str | None:
    """Return a stable, readable representation of an openpyxl color object."""
    if color is None:
        return None
    color_type = getattr(color, "type", None)
    if color_type == "rgb":
        rgb = getattr(color, "rgb", None)
        return str(rgb) if rgb else None
    if color_type == "indexed":
        value = getattr(color, "indexed", None)
        return f"indexed:{value}" if value is not None else None
    if color_type == "theme":
        value = getattr(color, "theme", None)
        tint = getattr(color, "tint", 0)
        return f"theme:{value}:tint:{tint}" if value is not None else None
    if color_type == "auto":
        return "auto"
    return str(color_type) if color_type else None


def _fill_color_text(cell: Cell) -> str | None:
    fill = cell.fill
    if fill is None or getattr(fill, "fill_type", None) is None:
        return None
    return _style_color_text(getattr(fill, "fgColor", None))


@dataclass(frozen=True)
class LoadedCell:
    """Value + useful display/style metadata copied from an Excel cell."""

    row: int
    column: int
    coordinate: str
    value: FillValue
    number_format: str | None
    font_name: str | None
    font_size: float | None
    bold: bool
    italic: bool
    underline: str | bool | None
    font_color: str | None
    fill_color: str | None
    fill_type: str | None
    horizontal: str | None
    vertical: str | None
    wrap_text: bool | None
    hyperlink: str | None
    style_id: int | None
    style_name: str | None
    comment: str | None


@dataclass(frozen=True)
class LoadedColumn:
    """A table column. Iterating yields LoadedCell objects; .values gives raw values."""

    index: int
    letter: str
    header: str
    header_cell: LoadedCell | None
    cells: list[LoadedCell]

    def __iter__(self):
        return iter(self.cells)

    def __len__(self) -> int:
        return len(self.cells)

    def __getitem__(self, index: int) -> LoadedCell:
        return self.cells[index]

    @property
    def values(self) -> list[FillValue]:
        return [cell.value for cell in self.cells]


@dataclass(frozen=True)
class LoadedRow:
    """A table/workbook row. Iterating yields LoadedCell objects; .values gives raw values."""

    row: int
    cells: list[LoadedCell]

    def __iter__(self):
        return iter(self.cells)

    def __len__(self) -> int:
        return len(self.cells)

    def __getitem__(self, index: int) -> LoadedCell:
        return self.cells[index]

    @property
    def values(self) -> list[FillValue]:
        return [cell.value for cell in self.cells]


class LoadedTable:
    """Read-only snapshot of one Excel worksheet as a simple headered table."""

    def __init__(
        self,
        *,
        path: Path | None,
        sheet_name: str,
        header_row: int,
        headers: list[str],
        header_cells: list[LoadedCell],
        data_rows: list[LoadedRow],
        all_rows: list[LoadedRow],
    ) -> None:
        self.path = path
        self.sheet_name = sheet_name
        self.header_row = header_row
        self.headers = headers
        self.header_cells = header_cells
        self.rows = data_rows
        self.all_rows = all_rows

    @staticmethod
    def _snapshot_cell(cell: Cell) -> LoadedCell:
        link = None
        if cell.hyperlink is not None:
            link = cell.hyperlink.target or cell.hyperlink.location or cell.hyperlink.display
        comment = cell.comment.text if cell.comment is not None else None
        return LoadedCell(
            row=cell.row,
            column=cell.column,
            coordinate=cell.coordinate,
            value=cell.value,
            number_format=cell.number_format,
            font_name=cell.font.name,
            font_size=float(cell.font.sz) if cell.font.sz is not None else None,
            bold=bool(cell.font.bold),
            italic=bool(cell.font.italic),
            underline=cell.font.underline,
            font_color=_style_color_text(cell.font.color),
            fill_color=_fill_color_text(cell),
            fill_type=cell.fill.fill_type,
            horizontal=cell.alignment.horizontal,
            vertical=cell.alignment.vertical,
            wrap_text=cell.alignment.wrap_text,
            hyperlink=link,
            style_id=getattr(cell, "style_id", None),
            style_name=cell.style,
            comment=comment,
        )

    @classmethod
    def from_worksheet(
        cls,
        ws: Worksheet,
        *,
        path: str | Path | None = None,
        header_row: int = 1,
    ) -> "LoadedTable":
        if header_row < 1:
            raise ValueError(f"header_row must be >= 1, got {header_row}")
        if header_row > ws.max_row:
            raise ValueError(f"header_row={header_row} is beyond sheet max_row={ws.max_row}")

        all_rows: list[LoadedRow] = []
        for row_idx in range(1, ws.max_row + 1):
            cells = [cls._snapshot_cell(ws.cell(row=row_idx, column=col_idx)) for col_idx in range(1, ws.max_column + 1)]
            all_rows.append(LoadedRow(row=row_idx, cells=cells))

        header_cells = all_rows[header_row - 1].cells if all_rows else []
        headers = ["" if cell.value is None else str(cell.value).strip() for cell in header_cells]
        data_rows = [row for row in all_rows if row.row > header_row]
        return cls(
            path=Path(path) if path is not None else None,
            sheet_name=ws.title,
            header_row=header_row,
            headers=headers,
            header_cells=header_cells,
            data_rows=data_rows,
            all_rows=all_rows,
        )

    def _header_matches(self, wanted: str) -> list[int]:
        exact = [idx for idx, header in enumerate(self.headers, start=1) if header == wanted]
        if exact:
            return exact
        folded = nodia_casefold(wanted)
        return [idx for idx, header in enumerate(self.headers, start=1) if nodia_casefold(header) == folded]

    @staticmethod
    def _is_excel_column_letters(text: str) -> bool:
        if not text or len(text) > 3:
            return False
        if not all("A" <= ch <= "Z" for ch in text.upper()):
            return False
        try:
            idx = column_to_index(text)
        except Exception:
            return False
        # Excel's last column is XFD = 16384.
        return 1 <= idx <= 16384

    def _resolve_column(self, column: ColumnRef) -> int:
        if isinstance(column, int):
            if column < 1 or column > len(self.headers):
                raise IndexError(f"Column index must be in 1..{len(self.headers)}, got {column}")
            return column
        if not isinstance(column, str):
            raise TypeError(f"Column must be int or str, got {type(column).__name__}")

        text = column.strip()
        if not text:
            raise ValueError("Column string cannot be empty")

        # Prefer header names over letters. This keeps get_column("CMU") useful
        # when the table really has a CMU header, even though CMU is also a
        # syntactically valid Excel column label.
        matches = self._header_matches(text)
        if len(matches) == 1:
            return matches[0]
        if len(matches) > 1:
            raise ValueError(
                f"Header {text!r} is ambiguous; matching columns: "
                + ", ".join(index_to_column(i) for i in matches)
                + ". Use a column number or letter instead."
            )

        if self._is_excel_column_letters(text):
            idx = column_to_index(text)
            if idx > len(self.headers):
                raise IndexError(f"Column {text!r} resolves to {idx}, but the table has only {len(self.headers)} columns")
            return idx

        raise KeyError(f"No column named {column!r}")

    def get_column(self, column: ColumnRef) -> LoadedColumn:
        idx = self._resolve_column(column)
        header = self.headers[idx - 1] if idx - 1 < len(self.headers) else ""
        header_cell = self.header_cells[idx - 1] if idx - 1 < len(self.header_cells) else None
        cells = [row.cells[idx - 1] for row in self.rows if idx - 1 < len(row.cells)]
        return LoadedColumn(
            index=idx,
            letter=index_to_column(idx),
            header=header,
            header_cell=header_cell,
            cells=cells,
        )

    def get_row(self, row: int) -> LoadedRow:
        if row < 1 or row > len(self.all_rows):
            raise IndexError(f"Row must be in 1..{len(self.all_rows)}, got {row}")
        return self.all_rows[row - 1]

    def column_values(self, column: ColumnRef, *, drop_blank: bool = False) -> list[FillValue]:
        values = self.get_column(column).values
        if not drop_blank:
            return values
        return [v for v in values if v is not None and not (isinstance(v, str) and not v.strip())]

    def to_records(self, *, drop_blank_rows: bool = True) -> list[dict[str, FillValue]]:
        records: list[dict[str, FillValue]] = []
        for row in self.rows:
            values = row.values
            if drop_blank_rows and all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
                continue
            rec: dict[str, FillValue] = {}
            for idx, header in enumerate(self.headers):
                if not header:
                    continue
                rec[header] = values[idx] if idx < len(values) else None
            records.append(rec)
        return records

class ExcelTable:
    """Small convenience wrapper around openpyxl for painless Excel tables.

    Designed for scripts where you want to write:

        from tables import ExcelTable
        tab = ExcelTable("scratch/out.xlsx")
        tab.set_cell("meow", column="A", row=4, bold=True)
        tab.save()

    Coordinates and sheet indexes are always 1-based, matching Excel.
    """

    def __init__(
        self,
        path: str | Path,
        *,
        sheet_name: str = "Sheet1",
        create: bool = True,
        overwrite: bool = False,
        autosave: bool = False,
    ) -> None:
        self.path = Path(path)
        self.autosave = autosave
        self._default_font_name: str | None = None
        self._default_font_size: int | float | None = None

        if self.path.exists() and not overwrite:
            self.wb = load_workbook(self.path)
            if sheet_name in self.wb.sheetnames:
                self.ws = self.wb[sheet_name]
            elif len(self.wb.worksheets) == 1 and self.wb.active.title in {"Sheet", "Sheet1"}:
                # Be friendly when opening a plain new workbook and caller wants a
                # specific first sheet name.
                self.ws = self.wb.active
                self._assert_sheet_name_available(sheet_name, allow_current=self.ws)
                self.ws.title = sheet_name
            else:
                self.ws = self.wb.create_sheet(sheet_name)
        else:
            if not create and not self.path.exists():
                raise FileNotFoundError(self.path)
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = sheet_name
        self._activate_ws(self.ws)

    def __enter__(self) -> "ExcelTable":
        return self

    def __exit__(self, exc_type: Any, exc: Any, tb: Any) -> None:
        if exc_type is None:
            self.save()

    def __repr__(self) -> str:
        return f"ExcelTable(path={str(self.path)!r}, active_sheet={self.active_sheet!s})"

    # ------------------------------------------------------------------
    # Sheet/workbook management
    # ------------------------------------------------------------------
    @property
    def sheet(self) -> Worksheet:
        """The current active worksheet object."""
        return self.ws

    @property
    def sheets(self) -> list[str]:
        """Worksheet names in workbook order."""
        return list(self.wb.sheetnames)

    @property
    def sheet_names(self) -> list[str]:
        return self.sheets

    @property
    def sheet_count(self) -> int:
        return len(self.wb.worksheets)

    @property
    def active_sheet(self) -> ActiveSheet:
        ws = self.ws
        return ActiveSheet(index=self.wb.worksheets.index(ws) + 1, name=ws.title)

    @property
    def max_row(self) -> int:
        return self.ws.max_row

    @property
    def max_column(self) -> int:
        return self.ws.max_column

    def _maybe_save(self) -> None:
        if self.autosave:
            self.save()

    def _activate_ws(self, ws: Worksheet) -> None:
        self.ws = ws
        self.wb.active = self.wb.worksheets.index(ws)

    def _assert_sheet_name_available(self, name: str, *, allow_current: Worksheet | None = None) -> None:
        if not isinstance(name, str) or not name.strip():
            raise ValueError("Sheet name must be a non-empty string")
        if name in self.wb.sheetnames and (allow_current is None or self.wb[name] is not allow_current):
            raise ValueError(f"Sheet {name!r} already exists")

    def _resolve_sheet(self, sheet: SheetRef = None) -> Worksheet:
        """Resolve None/name/1-based index/Worksheet to a worksheet.

        None means the current active sheet, so all old sheetless calls continue
        to work exactly as before.
        """
        if sheet is None:
            return self.ws
        if isinstance(sheet, Worksheet):
            if sheet not in self.wb.worksheets:
                raise ValueError("Worksheet does not belong to this workbook")
            return sheet
        if isinstance(sheet, int):
            if sheet < 1 or sheet > len(self.wb.worksheets):
                raise IndexError(f"Sheet index must be in 1..{len(self.wb.worksheets)}, got {sheet}")
            return self.wb.worksheets[sheet - 1]
        if isinstance(sheet, str):
            if sheet not in self.wb.sheetnames:
                raise KeyError(f"No sheet named {sheet!r}")
            return self.wb[sheet]
        raise TypeError(f"Sheet must be None, int, str, or Worksheet, got {type(sheet).__name__}")

    def _split_sheet_from_range(self, range_ref: str, sheet: SheetRef = None) -> tuple[Worksheet, str]:
        """Allow either sheet='Foo', range='A1:B2' or range='Foo!A1:B2'."""
        if "!" not in range_ref:
            return self._resolve_sheet(sheet), range_ref
        sheet_part, cell_range = range_ref.split("!", 1)
        if sheet is not None:
            raise ValueError("Specify sheet either via sheet=... or in the range string, not both")
        sheet_name = sheet_part.strip("'")
        return self._resolve_sheet(sheet_name), cell_range

    def set_active_sheet(self, sheet: SheetRef) -> "ExcelTable":
        ws = self._resolve_sheet(sheet)
        self._activate_ws(ws)
        self._maybe_save()
        return self

    def use_sheet(self, sheet: SheetRef, *, create: bool = True) -> "ExcelTable":
        """Backward-compatible alias for selecting a sheet.

        If passed a missing string and create=True, creates and activates it.
        """
        if isinstance(sheet, str) and sheet not in self.wb.sheetnames and create:
            self.add_sheet(sheet)
        else:
            self.set_active_sheet(sheet)
        return self

    def add_sheet(self, name: str) -> Worksheet:
        """Add a new worksheet and make it active."""
        self._assert_sheet_name_available(name)
        ws = self.wb.create_sheet(name)
        self._activate_ws(ws)
        self._maybe_save()
        return ws

    def create_sheet(self, sheet_name: str, *, use: bool = True) -> Worksheet:
        """Backward-compatible alias. Prefer add_sheet(name)."""
        self._assert_sheet_name_available(sheet_name)
        ws = self.wb.create_sheet(sheet_name)
        if use:
            self._activate_ws(ws)
        self._maybe_save()
        return ws

    def rename_sheet(self, sheet: SheetRef, new_name: str) -> Worksheet:
        ws = self._resolve_sheet(sheet)
        self._assert_sheet_name_available(new_name, allow_current=ws)
        ws.title = new_name
        if ws is self.ws:
            self._activate_ws(ws)
        self._maybe_save()
        return ws

    def remove_sheet(self, sheet: SheetRef) -> None:
        """Remove a sheet. Excel requires at least one sheet to remain."""
        ws = self._resolve_sheet(sheet)
        if len(self.wb.worksheets) <= 1:
            raise ValueError("Cannot remove the only sheet in a workbook")
        was_active = ws is self.ws
        self.wb.remove(ws)
        if was_active:
            self._activate_ws(self.wb.worksheets[0])
        self._maybe_save()

    def save(self, path: str | Path | None = None) -> Path:
        target = Path(path) if path is not None else self.path
        target.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(target)
        self.path = target
        return target

    def close(self) -> None:
        self.wb.close()

    def set_default_font(
        self,
        font_name: str,
        font_size: int | float | None = None,
        *,
        apply_existing: bool = True,
    ) -> "ExcelTable":
        """Set the workbook/helper default font for future writes.

        openpyxl does not have a reliable global workbook font switch that
        automatically affects every new cell written by a helper method. This
        method therefore stores a helper-level default used by ``set_cell()``
        whenever no explicit ``font_name``/``font_size`` is supplied.

        If ``apply_existing`` is true, it also updates fonts on cells that
        already exist in the workbook while preserving bold/italic/underline,
        color, and other font attributes.
        """
        clean_name = str(font_name or "").strip()
        if not clean_name:
            raise ValueError("font_name must be a non-empty string")

        if font_size is not None:
            try:
                font_size = float(font_size)
            except (TypeError, ValueError) as exc:
                raise ValueError(f"font_size must be numeric, got {font_size!r}") from exc

        self._default_font_name = clean_name
        self._default_font_size = font_size

        # Best-effort update of Excel's built-in Normal style. This helps with
        # cells that Excel creates visually, but the stored helper default above
        # is what guarantees future ExcelTable.set_cell() calls use the font.
        for style in getattr(self.wb, "_named_styles", []):  # noqa: SLF001 - openpyxl stores named styles here.
            if getattr(style, "name", None) == "Normal":
                old_font = copy(style.font)
                old_font.name = clean_name
                if font_size is not None:
                    old_font.sz = font_size
                style.font = old_font
                break

        if apply_existing:
            for ws in self.wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        old_font = copy(cell.font)
                        old_font.name = clean_name
                        if font_size is not None:
                            old_font.sz = font_size
                        cell.font = old_font

        self._maybe_save()
        return self

    @classmethod
    def load_table(
        cls,
        path: str | Path,
        *,
        sheet: SheetRef = None,
        header_row: int = 1,
        data_only: bool = True,
    ) -> LoadedTable:
        """Load one worksheet from an .xlsx file as a read-only LoadedTable.

        The returned object has convenient get_column()/get_row() helpers and
        preserves useful cell metadata such as font, fill, bold/italic,
        wrapping, number format, hyperlinks, comments, and style identifiers.
        """
        source = Path(path)
        wb = load_workbook(source, data_only=data_only)
        if sheet is None:
            ws = wb.active
        elif isinstance(sheet, Worksheet):
            ws = sheet
        elif isinstance(sheet, int):
            if sheet < 1 or sheet > len(wb.worksheets):
                raise IndexError(f"Sheet index must be in 1..{len(wb.worksheets)}, got {sheet}")
            ws = wb.worksheets[sheet - 1]
        elif isinstance(sheet, str):
            if sheet not in wb.sheetnames:
                raise KeyError(f"No sheet named {sheet!r}")
            ws = wb[sheet]
        else:
            raise TypeError(f"Sheet must be None, int, str, or Worksheet, got {type(sheet).__name__}")
        return LoadedTable.from_worksheet(ws, path=source, header_row=header_row)

    def to_loaded_table(self, *, sheet: SheetRef = None, header_row: int = 1) -> LoadedTable:
        """Snapshot an already-open worksheet as a LoadedTable."""
        ws = self._resolve_sheet(sheet)
        return LoadedTable.from_worksheet(ws, path=self.path, header_row=header_row)

    def get_column(self, column: ColumnRef, *, sheet: SheetRef = None, header_row: int = 1) -> LoadedColumn:
        """Convenience wrapper around to_loaded_table(...).get_column(...)."""
        return self.to_loaded_table(sheet=sheet, header_row=header_row).get_column(column)

    def get_row(self, row: int, *, sheet: SheetRef = None, header_row: int = 1) -> LoadedRow:
        """Convenience wrapper around to_loaded_table(...).get_row(...)."""
        return self.to_loaded_table(sheet=sheet, header_row=header_row).get_row(row)

    # ------------------------------------------------------------------
    # Coordinates
    # ------------------------------------------------------------------
    @staticmethod
    def col(column: ColumnRef) -> int:
        return column_to_index(column)

    @staticmethod
    def col_name(column: int) -> str:
        return index_to_column(column)

    @staticmethod
    def cell_name(column: ColumnRef, row: int) -> str:
        return cell_ref(column, row)

    def cell(self, column: ColumnRef, row: int, *, sheet: SheetRef = None) -> Cell:
        return self._resolve_sheet(sheet).cell(row=row, column=column_to_index(column))

    @staticmethod
    def _coerce_cell_value(content: Any) -> FillValue:
        """Convert friendly Python objects to values Excel can store.

        openpyxl only accepts scalar cell values. This helper keeps all normal
        Excel values unchanged, and converts custom metadata wrappers such as
        MetaIS CI/relation objects using str(obj), so calls like
        ``tab.set_cell(citype, ...)`` work as expected.
        """
        if content is None or isinstance(content, (str, int, float, bool, date, datetime, time, Decimal)):
            return content
        return str(content)

    # ------------------------------------------------------------------
    # Main API
    # ------------------------------------------------------------------
    def set_cell(
        self,
        content: RawFillValue = None,
        *,
        column: ColumnRef,
        row: int,
        sheet: SheetRef = None,
        format: str | None = "general",
        alignment: str | None = "right",
        bold: bool = False,
        italic: bool = False,
        underscore: bool = False,
        underline: bool | None = None,
        link: str | None = None,
        cell_color: ColorLike = None,
        text_color: ColorLike = "black",
        cell_border: Any = True,
        border_left: Any = None,
        border_right: Any = None,
        border_top: Any = None,
        border_bottom: Any = None,
        text_style: Any = None,
        wrap_text: bool | None = None,
        font_name: str | None = None,
        font_size: int | float | None = None,
        comment: str | None = None,
    ) -> Cell:
        """Set a single cell's value and style.

        `content=None` means the cell is intentionally blank but can still be
        colored/bordered/formatted. `sheet=None` means the active sheet.
        """
        if row < 1:
            raise ValueError(f"Row must be >= 1, got {row}")

        style_alias = parse_text_style(text_style)
        bold = bool(style_alias.get("bold", bold))
        italic = bool(style_alias.get("italic", italic))
        if underline is not None:
            underscore = underline
        underscore = bool(style_alias.get("underscore", underscore))
        text_color = style_alias.get("text_color", text_color)
        font_name = style_alias.get("font_name", font_name)
        font_size = style_alias.get("font_size", font_size)
        if font_name is None:
            font_name = self._default_font_name
        if font_size is None:
            font_size = self._default_font_size

        cell = self.cell(column, row, sheet=sheet)
        cell.value = self._coerce_cell_value(content)
        cell.number_format = normalize_number_format(format) or cell.number_format
        cell.alignment = make_alignment(alignment, wrap_text=wrap_text)
        cell.font = make_font(
            bold=bold,
            italic=italic,
            underscore=underscore,
            text_color=text_color,
            font_name=font_name,
            font_size=font_size,
        )
        cell.fill = make_fill(cell_color)
        cell.border = make_border(
            cell_border,
            border_left=border_left,
            border_right=border_right,
            border_top=border_top,
            border_bottom=border_bottom,
        )

        if link is not None:
            cell.hyperlink = link
            # Excel convention: hyperlinks are blue and underlined unless caller
            # explicitly asked for a different text style.
            if text_color == "black" and not underscore and text_style is None:
                cell.font = make_font(
                    bold=bold,
                    italic=italic,
                    underscore=True,
                    text_color="blue",
                    font_name=font_name,
                    font_size=font_size,
                )
        else:
            cell.hyperlink = None

        if comment is not None:
            from openpyxl.comments import Comment

            cell.comment = Comment(comment, "ExcelTable")

        self._maybe_save()
        return cell

    def set_row(
        self,
        row: int,
        values: Sequence[RawFillValue],
        *,
        start_column: ColumnRef = 1,
        sheet: SheetRef = None,
        **style: Any,
    ) -> list[Cell]:
        col = column_to_index(start_column)
        cells: list[Cell] = []
        for offset, value in enumerate(values):
            cells.append(self.set_cell(value, column=col + offset, row=row, sheet=sheet, **style))
        return cells

    def append_row(self, values: Sequence[RawFillValue], *, sheet: SheetRef = None, **style: Any) -> list[Cell]:
        ws = self._resolve_sheet(sheet)
        row = ws.max_row + 1 if ws.max_row > 1 or ws.cell(1, 1).value is not None else 1
        return self.set_row(row, values, sheet=ws, **style)

    def set_header(
        self,
        values_or_row: Sequence[RawFillValue] | int,
        values: Sequence[RawFillValue] | None = None,
        *,
        row: int | None = None,
        start_column: ColumnRef = 1,
        sheet: SheetRef = None,
        cell_color: ColorLike = "#D9EAD3",
        text_color: ColorLike = "black",
        **style: Any,
    ) -> list[Cell]:
        """Write a styled header row.

        Preferred usage:

            tab.set_header(["A", "B", "C"])          # row 1
            tab.set_header(["A", "B", "C"], row=2)  # explicit row

        Backward-compatible usage is still accepted:

            tab.set_header(1, ["A", "B", "C"])
        """
        if isinstance(values_or_row, int):
            if values is None:
                raise TypeError("set_header(row, values) requires values")
            if row is not None and row != values_or_row:
                raise ValueError("Specify the header row either positionally or with row=..., not both")
            header_row = values_or_row
            header_values = values
        else:
            if values is not None:
                raise TypeError("When the first argument is values, do not pass a second positional values argument")
            header_row = 1 if row is None else row
            header_values = values_or_row

        defaults = {
            "bold": True,
            "alignment": "center",
            "cell_color": cell_color,
            "text_color": text_color,
            "wrap_text": True,
        }
        defaults.update(style)
        return self.set_row(header_row, header_values, start_column=start_column, sheet=sheet, **defaults)

    # ------------------------------------------------------------------
    # Color/style convenience helpers
    # ------------------------------------------------------------------
    def set_cell_color(self, *, column: ColumnRef, row: int, color: ColorLike, sheet: SheetRef = None) -> Cell:
        cell = self.cell(column, row, sheet=sheet)
        cell.fill = make_fill(color)
        self._maybe_save()
        return cell

    def set_text_color(self, *, column: ColumnRef, row: int, color: ColorLike, sheet: SheetRef = None) -> Cell:
        cell = self.cell(column, row, sheet=sheet)
        old = copy(cell.font)
        old.color = make_font(text_color=color).color
        cell.font = old
        self._maybe_save()
        return cell

    def set_row_color(
        self,
        row: int,
        color: ColorLike,
        *,
        from_column: ColumnRef = 1,
        to_column: ColumnRef | None = None,
        sheet: SheetRef = None,
    ) -> None:
        ws = self._resolve_sheet(sheet)
        start = column_to_index(from_column)
        end = column_to_index(to_column) if to_column is not None else ws.max_column
        fill = make_fill(color)
        for col in range(start, end + 1):
            ws.cell(row=row, column=col).fill = copy(fill)
        self._maybe_save()

    def set_column_color(
        self,
        column: ColumnRef,
        color: ColorLike,
        *,
        from_row: int = 1,
        to_row: int | None = None,
        sheet: SheetRef = None,
    ) -> None:
        ws = self._resolve_sheet(sheet)
        col = column_to_index(column)
        end = to_row if to_row is not None else ws.max_row
        fill = make_fill(color)
        for row in range(from_row, end + 1):
            ws.cell(row=row, column=col).fill = copy(fill)
        self._maybe_save()

    def set_range_color(self, range_ref: str, color: ColorLike, *, sheet: SheetRef = None) -> None:
        ws, clean_range = self._split_sheet_from_range(range_ref, sheet)
        min_col, min_row, max_col, max_row = range_boundaries(clean_range)
        fill = make_fill(color)
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                ws.cell(row=row, column=col).fill = copy(fill)
        self._maybe_save()

    def set_border(
        self,
        range_ref: str,
        *,
        sheet: SheetRef = None,
        cell_border: Any = True,
        border_left: Any = None,
        border_right: Any = None,
        border_top: Any = None,
        border_bottom: Any = None,
    ) -> None:
        ws, clean_range = self._split_sheet_from_range(range_ref, sheet)
        border = make_border(
            cell_border,
            border_left=border_left,
            border_right=border_right,
            border_top=border_top,
            border_bottom=border_bottom,
        )
        min_col, min_row, max_col, max_row = range_boundaries(clean_range)
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                ws.cell(row=row, column=col).border = copy(border)
        self._maybe_save()

    # ------------------------------------------------------------------
    # Structure helpers
    # ------------------------------------------------------------------
    def _shift_hyperlinks_for_row_insert(self, ws: Worksheet, insert_at: int, amount: int) -> list[tuple[int, int, Any]]:
        shifted: list[tuple[int, int, Any]] = []
        for cell in list(ws._cells.values()):  # noqa: SLF001 - openpyxl stores links on cells.
            if cell.hyperlink is None:
                continue
            new_row = cell.row + amount if cell.row >= insert_at else cell.row
            new_link = copy(cell.hyperlink)
            new_link.ref = f"{get_column_letter(cell.column)}{new_row}"
            shifted.append((new_row, cell.column, new_link))
        return shifted

    def _shift_hyperlinks_for_col_insert(self, ws: Worksheet, insert_at: int, amount: int) -> list[tuple[int, int, Any]]:
        shifted: list[tuple[int, int, Any]] = []
        for cell in list(ws._cells.values()):  # noqa: SLF001 - openpyxl stores links on cells.
            if cell.hyperlink is None:
                continue
            new_col = cell.column + amount if cell.column >= insert_at else cell.column
            new_link = copy(cell.hyperlink)
            new_link.ref = f"{get_column_letter(new_col)}{cell.row}"
            shifted.append((cell.row, new_col, new_link))
        return shifted

    def _restore_hyperlinks(self, ws: Worksheet, links: list[tuple[int, int, Any]]) -> None:
        for row, col, link in links:
            ws.cell(row=row, column=col).hyperlink = link

    def insert_row(self, row: int, amount: int = 1, *, sheet: SheetRef = None) -> None:
        """Insert empty row(s) before `row` and shift existing cells/styles/links."""
        if row < 1:
            raise ValueError(f"Row must be >= 1, got {row}")
        if amount < 1:
            raise ValueError(f"Amount must be >= 1, got {amount}")
        ws = self._resolve_sheet(sheet)
        links = self._shift_hyperlinks_for_row_insert(ws, row, amount)
        ws.insert_rows(row, amount)
        self._restore_hyperlinks(ws, links)
        self._maybe_save()

    def insert_column(self, column: ColumnRef, amount: int = 1, *, sheet: SheetRef = None) -> None:
        """Insert empty column(s) before `column` and shift existing cells/styles/links."""
        if amount < 1:
            raise ValueError(f"Amount must be >= 1, got {amount}")
        ws = self._resolve_sheet(sheet)
        col = column_to_index(column)
        links = self._shift_hyperlinks_for_col_insert(ws, col, amount)
        ws.insert_cols(col, amount)
        self._restore_hyperlinks(ws, links)
        self._maybe_save()

    def delete_row(self, row: int, amount: int = 1, *, sheet: SheetRef = None) -> None:
        self._resolve_sheet(sheet).delete_rows(row, amount)
        self._maybe_save()

    def delete_column(self, column: ColumnRef, amount: int = 1, *, sheet: SheetRef = None) -> None:
        self._resolve_sheet(sheet).delete_cols(column_to_index(column), amount)
        self._maybe_save()

    def freeze_panes(self, cell: str | None = "A2", *, sheet: SheetRef = None) -> None:
        self._resolve_sheet(sheet).freeze_panes = cell
        self._maybe_save()

    def _default_filter_range(
        self,
        ws: Worksheet,
        *,
        header_row: int = 1,
        from_column: ColumnRef = 1,
        to_column: ColumnRef | None = None,
        to_row: int | None = None,
    ) -> str:
        if header_row < 1:
            raise ValueError(f"header_row must be >= 1, got {header_row}")

        min_col = column_to_index(from_column)
        max_col = column_to_index(to_column) if to_column is not None else ws.max_column
        max_row = to_row if to_row is not None else ws.max_row

        if min_col > max_col:
            raise ValueError("from_column must be <= to_column")
        if max_row < header_row:
            max_row = header_row

        return (
            f"{get_column_letter(min_col)}{header_row}:"
            f"{get_column_letter(max_col)}{max_row}"
        )

    def enable_filter(
        self,
        range_ref: str | None = None,
        *,
        sheet: SheetRef = None,
        header_row: int = 1,
        from_column: ColumnRef = 1,
        to_column: ColumnRef | None = None,
        to_row: int | None = None,
    ) -> None:
        """Enable Excel's clickable Data > Filter dropdowns for a range.

        Typical use after writing a normal table with headers in row 1:

            tab.enable_filter()

        By default this applies the filter to the used range of the selected
        sheet, starting at row 1.  You can also pass an explicit Excel range:

            tab.enable_filter("A1:D100")
            tab.enable_filter("Entities!A1:F200")

        Or build the range from parts:

            tab.enable_filter(sheet="Entities", from_column="A", to_column="F")
        """
        if range_ref is None:
            ws = self._resolve_sheet(sheet)
            clean_range = self._default_filter_range(
                ws,
                header_row=header_row,
                from_column=from_column,
                to_column=to_column,
                to_row=to_row,
            )
        else:
            ws, clean_range = self._split_sheet_from_range(range_ref, sheet)

        ws.auto_filter.ref = clean_range
        self._maybe_save()

    def disable_filter(self, *, sheet: SheetRef = None) -> None:
        """Disable Excel's clickable filter dropdowns on a sheet."""
        ws = self._resolve_sheet(sheet)
        ws.auto_filter.ref = None
        ws.auto_filter.filterColumn = []
        ws.auto_filter.sortState = None
        self._maybe_save()

    def toggle_filter(
        self,
        range_ref: str | None = None,
        *,
        sheet: SheetRef = None,
        header_row: int = 1,
        from_column: ColumnRef = 1,
        to_column: ColumnRef | None = None,
        to_row: int | None = None,
    ) -> None:
        """Toggle Excel's filter dropdowns for a sheet."""
        ws = self._resolve_sheet(sheet)
        if ws.auto_filter.ref:
            self.disable_filter(sheet=ws)
        else:
            self.enable_filter(
                range_ref,
                sheet=ws,
                header_row=header_row,
                from_column=from_column,
                to_column=to_column,
                to_row=to_row,
            )

    def autofilter(
        self,
        range_ref: str | None = None,
        *,
        sheet: SheetRef = None,
        header_row: int = 1,
        from_column: ColumnRef = 1,
        to_column: ColumnRef | None = None,
        to_row: int | None = None,
    ) -> None:
        """Backward-compatible alias for enable_filter()."""
        self.enable_filter(
            range_ref,
            sheet=sheet,
            header_row=header_row,
            from_column=from_column,
            to_column=to_column,
            to_row=to_row,
        )

    def merge_cells(
        self,
        *,
        start_row: int,
        start_column: ColumnRef,
        end_row: int,
        end_column: ColumnRef,
        sheet: SheetRef = None,
        content: RawFillValue = None,
        **style: Any,
    ) -> Cell:
        ws = self._resolve_sheet(sheet)
        start_col = column_to_index(start_column)
        end_col = column_to_index(end_column)
        ws.merge_cells(start_row=start_row, start_column=start_col, end_row=end_row, end_column=end_col)
        return self.set_cell(content, row=start_row, column=start_col, sheet=ws, **style)

    def set_width(self, column: ColumnRef, width: float, *, sheet: SheetRef = None) -> None:
        ws = self._resolve_sheet(sheet)
        ws.column_dimensions[index_to_column(column_to_index(column))].width = width
        self._maybe_save()

    def set_height(self, row: int, height: float, *, sheet: SheetRef = None) -> None:
        self._resolve_sheet(sheet).row_dimensions[row].height = height
        self._maybe_save()

    def wrap_column(
        self,
        column: ColumnRef,
        *,
        from_row: int = 1,
        to_row: int | None = None,
        sheet: SheetRef = None,
        vertical: str | None = "top",
    ) -> None:
        """Enable wrapped text for one column without changing other cell styles."""
        if from_row < 1:
            raise ValueError(f"from_row must be >= 1, got {from_row}")
        ws = self._resolve_sheet(sheet)
        col = column_to_index(column)
        end = to_row if to_row is not None else ws.max_row
        for row in range(from_row, end + 1):
            cell = ws.cell(row=row, column=col)
            old = copy(cell.alignment)
            old.wrap_text = True
            if vertical is not None:
                old.vertical = vertical
            cell.alignment = old
        self._maybe_save()

    def wrap_columns(
        self,
        columns: Sequence[ColumnRef],
        *,
        from_row: int = 1,
        to_row: int | None = None,
        sheet: SheetRef = None,
        vertical: str | None = "top",
    ) -> None:
        """Enable wrapped text for multiple columns."""
        for column in columns:
            self.wrap_column(column, from_row=from_row, to_row=to_row, sheet=sheet, vertical=vertical)

    @classmethod
    def _cell_text_fragment_width_units(
        cls,
        cell: Cell,
        text: str,
        *,
        font_aware: bool = True,
    ) -> float:
        """Return display-width units for a text fragment in one cell.

        This uses the same font-aware width heuristic as ``autofit()``. It is
        intentionally value-preserving: callers can estimate wrapped line counts
        without inserting hard ``\n`` line breaks into the workbook.
        """
        width = cls._line_text_width_units(text)
        if not font_aware:
            return width
        font = cell.font
        return width * cls._font_width_scale(
            getattr(font, "name", None),
            getattr(font, "sz", None),
            bold=bool(getattr(font, "bold", False)),
        )

    @classmethod
    def _estimated_wrapped_line_count(
        cls,
        cell: Cell,
        *,
        max_line_width: float,
        font_aware: bool = True,
    ) -> int:
        """Estimate how many visual lines Excel word-wrap will need.

        Unlike the old implementation, this does not use ``ceil(len(text) / N)``.
        It greedily packs whitespace-separated tokens into visual lines according
        to the current column width. Existing hard newlines are honored, but no
        hard newlines are written back to the cell value.
        """
        value = cell.value
        if value is None:
            return 1
        text = str(value)
        if not text:
            return 1

        max_line_width = max(1.0, float(max_line_width))
        total_lines = 0
        space_width = cls._cell_text_fragment_width_units(cell, " ", font_aware=font_aware)

        for paragraph in text.replace("\u00a0", " ").splitlines() or [""]:
            words = paragraph.split()
            if not words:
                total_lines += 1
                continue

            line_width = 0.0
            for word in words:
                word_width = cls._cell_text_fragment_width_units(cell, word, font_aware=font_aware)

                # Very long unbroken tokens, such as UUIDs or technical names,
                # still need intra-token wrapping. Account for that without
                # changing the cell's text.
                if word_width > max_line_width:
                    if line_width > 0:
                        total_lines += 1
                        line_width = 0.0
                    full_lines = int(word_width // max_line_width)
                    remainder = word_width - (full_lines * max_line_width)
                    if remainder <= 1e-9:
                        total_lines += max(1, full_lines)
                        line_width = 0.0
                    else:
                        total_lines += full_lines
                        line_width = remainder
                    continue

                candidate_width = word_width if line_width == 0 else line_width + space_width + word_width
                if candidate_width <= max_line_width:
                    line_width = candidate_width
                else:
                    total_lines += 1
                    line_width = word_width

            if line_width > 0:
                total_lines += 1

        return max(1, total_lines)

    def autofit_row_heights(
        self,
        *,
        sheet: SheetRef = None,
        from_row: int = 1,
        to_row: int | None = None,
        columns: Sequence[ColumnRef] | None = None,
        min_height: float = 15,
        max_height: float = 120,
        line_height: float = 15,
        width_to_chars_factor: float = 1.0,
        font_aware: bool = True,
    ) -> None:
        """Approximate wrapped-text row heights for selected columns.

        Excel calculates true auto-height only in the GUI. This helper keeps the
        original cell text intact, estimates where Excel-style word wrapping will
        create visual lines for the current column widths, and sets row heights
        from the resulting line counts.

        ``width_to_chars_factor`` is kept for backward compatibility; it now
        scales the available column-width units rather than switching back to
        crude character-count wrapping.
        """
        if from_row < 1:
            raise ValueError(f"from_row must be >= 1, got {from_row}")
        ws = self._resolve_sheet(sheet)
        end_row = to_row if to_row is not None else ws.max_row
        if end_row < from_row:
            return

        if columns is None:
            col_indices = list(range(1, ws.max_column + 1))
        else:
            col_indices = [column_to_index(c) for c in columns]

        for row in range(from_row, end_row + 1):
            max_lines = 1
            for col in col_indices:
                letter = get_column_letter(col)
                width = ws.column_dimensions[letter].width or 10
                available_width = max(1.0, float(width) * float(width_to_chars_factor))
                cell = ws.cell(row=row, column=col)
                max_lines = max(
                    max_lines,
                    self._estimated_wrapped_line_count(
                        cell,
                        max_line_width=available_width,
                        font_aware=font_aware,
                    ),
                )
            height = max(min_height, min(max_height, max_lines * line_height))
            ws.row_dimensions[row].height = height
        self._maybe_save()

    # ------------------------------------------------------------------
    # Sorting helpers
    # ------------------------------------------------------------------
    @staticmethod
    def _is_blank_sort_value(value: Any) -> bool:
        return value is None or value == ""

    @staticmethod
    def _sort_key_value(value: Any, *, case_sensitive: bool = False) -> tuple[int, Any]:
        """Return a predictable key for Excel-ish mixed values.

        Excel columns are usually homogeneous, but report scripts often mix
        strings, booleans, dates, and the occasional ``None``.  This keeps the
        sort stable and deterministic instead of letting Python raise on mixed
        incomparable types.
        """
        if isinstance(value, bool):
            return (0, int(value))
        if isinstance(value, (int, float, Decimal)):
            return (1, value)
        if isinstance(value, (datetime, date, time)):
            return (2, value.isoformat())
        if isinstance(value, str):
            return (3, value if case_sensitive else value.casefold())
        return (4, str(value))

    @staticmethod
    def _snapshot_cell(cell: Cell) -> dict[str, Any]:
        return {
            "value": cell.value,
            "style": copy(cell._style),  # noqa: SLF001 - openpyxl stores full style here.
            "number_format": cell.number_format,
            "font": copy(cell.font),
            "fill": copy(cell.fill),
            "border": copy(cell.border),
            "alignment": copy(cell.alignment),
            "protection": copy(cell.protection),
            "hyperlink": copy(cell.hyperlink) if cell.hyperlink is not None else None,
            "comment": copy(cell.comment) if cell.comment is not None else None,
        }

    @staticmethod
    def _restore_cell(cell: Cell, snap: dict[str, Any]) -> None:
        cell.value = snap["value"]
        cell._style = copy(snap["style"])  # noqa: SLF001 - restoring exact copied style.
        cell.number_format = snap["number_format"]
        cell.font = copy(snap["font"])
        cell.fill = copy(snap["fill"])
        cell.border = copy(snap["border"])
        cell.alignment = copy(snap["alignment"])
        cell.protection = copy(snap["protection"])
        cell.hyperlink = copy(snap["hyperlink"]) if snap["hyperlink"] is not None else None
        cell.comment = copy(snap["comment"]) if snap["comment"] is not None else None

    def _assert_sort_range_has_no_merged_cells(
        self,
        ws: Worksheet,
        *,
        min_col: int,
        max_col: int,
        min_row: int,
        max_row: int,
    ) -> None:
        for merged in ws.merged_cells.ranges:
            if (
                merged.min_col <= max_col
                and merged.max_col >= min_col
                and merged.min_row <= max_row
                and merged.max_row >= min_row
            ):
                raise ValueError(
                    "Cannot sort a range that intersects merged cells "
                    f"({merged.coord}). Unmerge first or choose a smaller range."
                )

    def sort_by_column(
        self,
        column: ColumnRef,
        *,
        sheet: SheetRef = None,
        header: bool = True,
        header_row: int = 1,
        from_row: int | None = None,
        to_row: int | None = None,
        from_column: ColumnRef = 1,
        to_column: ColumnRef | None = None,
        descending: bool = False,
        blank_last: bool = True,
        case_sensitive: bool = False,
    ) -> None:
        """Sort a rectangular table by one column while keeping rows together.

        Typical use with a header in row 1:

            tab.sort_by_column("C")

        This sorts rows 2..max_row by column C and moves values, cell styles,
        hyperlinks, comments, and row heights with their rows.  By default the
        sorted row width is columns 1..max_column of the selected sheet.

        Use ``from_column`` / ``to_column`` when the sortable table is only part
        of the sheet.  ``column`` must be inside that column range.
        """
        ws = self._resolve_sheet(sheet)

        sort_col = column_to_index(column)
        min_col = column_to_index(from_column)
        max_col = column_to_index(to_column) if to_column is not None else ws.max_column

        if header_row < 1:
            raise ValueError(f"header_row must be >= 1, got {header_row}")
        if min_col > max_col:
            raise ValueError("from_column must be <= to_column")
        if sort_col < min_col or sort_col > max_col:
            raise ValueError(
                f"Sort column {index_to_column(sort_col)} is outside "
                f"the sorted table range {index_to_column(min_col)}:{index_to_column(max_col)}"
            )

        start_row = from_row if from_row is not None else (header_row + 1 if header else header_row)
        end_row = to_row if to_row is not None else ws.max_row

        if start_row < 1:
            raise ValueError(f"from_row/start row must be >= 1, got {start_row}")
        if end_row < start_row:
            return

        self._assert_sort_range_has_no_merged_cells(
            ws, min_col=min_col, max_col=max_col, min_row=start_row, max_row=end_row
        )

        row_blocks: list[dict[str, Any]] = []
        for source_row in range(start_row, end_row + 1):
            sort_value = ws.cell(row=source_row, column=sort_col).value
            row_blocks.append(
                {
                    "source_row": source_row,
                    "blank": self._is_blank_sort_value(sort_value),
                    "key": self._sort_key_value(sort_value, case_sensitive=case_sensitive),
                    "height": ws.row_dimensions[source_row].height,
                    "hidden": ws.row_dimensions[source_row].hidden,
                    "outlineLevel": ws.row_dimensions[source_row].outlineLevel,
                    "collapsed": ws.row_dimensions[source_row].collapsed,
                    "cells": [
                        self._snapshot_cell(ws.cell(row=source_row, column=col))
                        for col in range(min_col, max_col + 1)
                    ],
                }
            )

        nonblank = [block for block in row_blocks if not block["blank"]]
        blank = [block for block in row_blocks if block["blank"]]
        nonblank.sort(key=lambda block: block["key"], reverse=descending)
        sorted_blocks = nonblank + blank if blank_last else blank + nonblank

        for target_row, block in zip(range(start_row, end_row + 1), sorted_blocks, strict=True):
            for offset, snap in enumerate(block["cells"]):
                cell = ws.cell(row=target_row, column=min_col + offset)
                self._restore_cell(cell, snap)
                if cell.hyperlink is not None:
                    cell.hyperlink.ref = cell.coordinate
            ws.row_dimensions[target_row].height = block["height"]
            ws.row_dimensions[target_row].hidden = block["hidden"]
            ws.row_dimensions[target_row].outlineLevel = block["outlineLevel"]
            ws.row_dimensions[target_row].collapsed = block["collapsed"]

        self._maybe_save()

    def sort_by_header(
        self,
        header_name: str,
        *,
        sheet: SheetRef = None,
        header_row: int = 1,
        from_column: ColumnRef = 1,
        to_column: ColumnRef | None = None,
        exact: bool = True,
        case_sensitive: bool = False,
        occurrence: int | None = None,
        **sort_options: Any,
    ) -> None:
        """Sort by the column whose header cell contains ``header_name``.

        Duplicate headers are allowed in Excel and in this module.  If a header
        name matches more than one column, either sort by an explicit column::

            tab.sort_by_column("C")

        or choose the 1-based matching occurrence::

            tab.sort_by_header("Valid", occurrence=2)

        Examples:

            tab.sort_by_header("Contains valid")
            tab.sort_by_header("Citype", descending=True)
        """
        ws = self._resolve_sheet(sheet)
        min_col = column_to_index(from_column)
        max_col = column_to_index(to_column) if to_column is not None else ws.max_column

        needle = header_name if case_sensitive else header_name.casefold()
        matches: list[int] = []
        for col in range(min_col, max_col + 1):
            value = ws.cell(row=header_row, column=col).value
            if value is None:
                continue
            haystack = str(value) if case_sensitive else str(value).casefold()
            if (haystack == needle) if exact else (needle in haystack):
                matches.append(col)

        if not matches:
            raise KeyError(f"No header matching {header_name!r} found in row {header_row}")

        if occurrence is not None:
            if occurrence < 1:
                raise ValueError(f"occurrence must be >= 1, got {occurrence}")
            if occurrence > len(matches):
                labels = ", ".join(index_to_column(col) for col in matches)
                raise ValueError(
                    f"Header {header_name!r} has only {len(matches)} matching column(s): {labels}; "
                    f"cannot use occurrence={occurrence}"
                )
            sort_col = matches[occurrence - 1]
        elif len(matches) > 1:
            labels = ", ".join(index_to_column(col) for col in matches)
            raise ValueError(
                f"Header {header_name!r} matched multiple columns: {labels}. "
                "Use sort_by_column('A') / sort_by_column(1), or pass "
                "sort_by_header(..., occurrence=N)."
            )
        else:
            sort_col = matches[0]

        self.sort_by_column(
            sort_col,
            sheet=ws,
            header=True,
            header_row=header_row,
            from_column=min_col,
            to_column=max_col,
            case_sensitive=case_sensitive,
            **sort_options,
        )


    @staticmethod
    def _line_text_width_units(text: str) -> float:
        """Estimate Excel width units for one visible line of text.

        Excel's real AutoFit uses rendered font metrics. openpyxl cannot render
        fonts, so this uses a conservative character-width heuristic. It is
        intentionally a bit generous: a slightly wide generated column is much
        better than clipped text.
        """
        total = 0.0
        for ch in str(text):
            code = ord(ch)
            if ch in "ilI.,:;|'`![](){} ":
                total += 0.55
            elif ch in "mwMW@#%&0123456789":
                total += 1.20
            elif ch == "\t":
                total += 4.0
            elif code >= 0x2E80:  # CJK / wide unicode blocks, rough fallback.
                total += 2.0
            else:
                total += 1.0
        return total

    @staticmethod
    def _font_width_scale(font_name: Any, font_size: Any, *, bold: bool = False) -> float:
        """Return a conservative scale from rendered text to Excel width units."""
        name = str(font_name or "").strip().casefold()
        if name in {"consolas", "courier new", "courier", "lucida console"}:
            scale = 1.23
        elif name in {"arial", "aptos", "calibri", "segoe ui"}:
            scale = 1.08
        elif name:
            scale = 1.15
        else:
            scale = 1.10

        try:
            size = float(font_size) if font_size is not None else 11.0
        except (TypeError, ValueError):
            size = 11.0
        scale *= max(0.75, min(1.60, size / 11.0))
        if bold:
            scale *= 1.05
        return scale

    @classmethod
    def _cell_text_width_units(cls, value: Any) -> float:
        """Approximate width contribution of a plain cell value.

        Multi-line text sizes to the longest visible line, not to the total
        string length. This method is kept for callers/tests that used the old
        helper directly; ``autofit()`` uses the font-aware cell variant below.
        """
        if value is None:
            return 0.0
        lines = str(value).splitlines() or [""]
        return float(max(cls._line_text_width_units(part) for part in lines))

    @classmethod
    def _cell_display_width_units(cls, cell: Cell, *, font_aware: bool = True) -> float:
        value_width = cls._cell_text_width_units(cell.value)
        if not font_aware:
            return value_width
        font = cell.font
        return value_width * cls._font_width_scale(
            getattr(font, "name", None),
            getattr(font, "sz", None),
            bold=bool(getattr(font, "bold", False)),
        )

    @staticmethod
    def _filter_header_cells(ws: Worksheet) -> set[tuple[int, int]]:
        """Return (row, column) cells that receive Excel filter arrows."""
        ref = ws.auto_filter.ref
        if not ref:
            return set()
        try:
            min_col, min_row, max_col, _max_row = range_boundaries(ref)
        except ValueError:
            return set()
        return {(min_row, col) for col in range(min_col, max_col + 1)}

    def autofit(
        self,
        *,
        sheet: SheetRef = None,
        min_width: float = 8,
        max_width: float | None = None,
        padding: float = 2.0,
        width_scale: float = 1.0,
        from_column: ColumnRef = 1,
        to_column: ColumnRef | None = None,
        respect_filter_dropdown: bool = True,
        filter_dropdown_padding: float = 6.0,
        font_aware: bool = True,
    ) -> None:
        """Set worksheet column widths from visible cell contents.

        This is a practical approximation of Excel AutoFit. Excel itself uses
        rendered font metrics and display/DPI details that openpyxl does not
        expose. The defaults therefore include a little padding and font-aware
        scaling so freshly opened workbooks are not clipped.
        """
        ws = self._resolve_sheet(sheet)
        start = column_to_index(from_column)
        end = column_to_index(to_column) if to_column is not None else ws.max_column
        filter_header_cells = self._filter_header_cells(ws) if respect_filter_dropdown else set()

        for col in range(start, end + 1):
            letter = get_column_letter(col)
            max_len = 0.0
            for cell in ws[letter]:
                value_width = self._cell_display_width_units(cell, font_aware=font_aware)
                # The filter dropdown only covers the header cell. So account
                # for it only in that header cell's candidate width. If a body
                # row is already wider than the header, the dropdown should not
                # add extra unnecessary width to the whole column.
                if (cell.row, cell.column) in filter_header_cells:
                    value_width += filter_dropdown_padding
                max_len = max(max_len, value_width)

            width = (max_len * float(width_scale)) + float(padding)
            if max_width is not None:
                width = min(max_width, width)
            width = max(min_width, width)
            ws.column_dimensions[letter].width = width
        self._maybe_save()

    def add_named_table(
        self,
        range_ref: str,
        *,
        sheet: SheetRef = None,
        name: str = "Table1",
        style: str = "TableStyleMedium2",
    ) -> None:
        from openpyxl.worksheet.table import Table, TableStyleInfo

        ws, clean_range = self._split_sheet_from_range(range_ref, sheet)
        table = Table(displayName=name, ref=clean_range)
        table.tableStyleInfo = TableStyleInfo(
            name=style,
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
        self._maybe_save()
