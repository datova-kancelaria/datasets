from __future__ import annotations

from dataclasses import dataclass
from typing import Any, Iterable, Mapping

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .colors import ColorLike, to_argb

BORDER_STYLES = {
    "dashdot",
    "dashdotdot",
    "dashed",
    "dotted",
    "double",
    "hair",
    "medium",
    "mediumdashdot",
    "mediumdashdotdot",
    "mediumdashed",
    "slantdashdot",
    "thick",
    "thin",
}

NUMBER_FORMATS: dict[str, str] = {
    "general": "General",
    "text": "@",
    "string": "@",
    "integer": "0",
    "int": "0",
    "number": "0.00",
    "float": "0.00",
    "decimal": "0.00",
    "percent": "0.00%",
    "percentage": "0.00%",
    "currency": '#,##0.00 €',
    "eur": '#,##0.00 €',
    "euro": '#,##0.00 €',
    "usd": '$#,##0.00',
    "date": "yyyy-mm-dd",
    "datetime": "yyyy-mm-dd hh:mm:ss",
    "time": "hh:mm:ss",
    "bool": "General",
    "boolean": "General",
}

DEFAULT_BORDER_COLOR = "#C6E0B4"  # thin light green, as requested.


def normalize_number_format(fmt: str | None) -> str | None:
    if fmt is None:
        return None
    key = fmt.strip().lower()
    return NUMBER_FORMATS.get(key, fmt)


def make_fill(color: ColorLike) -> PatternFill:
    argb = to_argb(color)
    if argb is None:
        return PatternFill(fill_type=None)
    return PatternFill(fill_type="solid", fgColor=argb)


def _is_border_style(value: Any) -> bool:
    return isinstance(value, str) and value.strip().lower().replace("-", "") in BORDER_STYLES


def _normalize_border_style(value: str) -> str:
    style = value.strip().lower().replace("-", "")
    if style not in BORDER_STYLES:
        raise ValueError(f"Unknown border style {value!r}")
    return style


def make_side(spec: Any = None, *, default_style: str = "thin", default_color: ColorLike = "black") -> Side:
    """Build an openpyxl Side from a flexible border spec.

    Supported specs:
      - None / False: no side
      - 'thin', 'dashed', ...: side style with default color
      - ('thin', 'black') or ('thin', '#000000')
      - ('thin', 'black', 'dashed'): permissive; the last style token wins
      - {'style': 'dashed', 'color': 'black'}
    """
    if spec is None or spec is False:
        return Side(style=None)
    if spec is True:
        spec = default_style

    style: str | None = default_style
    color: ColorLike = default_color

    if isinstance(spec, Mapping):
        style_val = spec.get("style", default_style)
        color_val = spec.get("color", default_color)
        style = None if style_val is None else _normalize_border_style(str(style_val))
        color = color_val  # type: ignore[assignment]
    elif isinstance(spec, str):
        if _is_border_style(spec):
            style = _normalize_border_style(spec)
        else:
            color = spec
    elif isinstance(spec, (tuple, list)):
        for item in spec:
            if item is None:
                continue
            if _is_border_style(item):
                style = _normalize_border_style(str(item))
            else:
                color = item  # type: ignore[assignment]
    else:
        raise TypeError(f"Unsupported border side spec: {spec!r}")

    if style is None:
        return Side(style=None)
    argb = to_argb(color, none_ok=False)
    return Side(style=style, color=argb)


def make_border(
    cell_border: Any = True,
    *,
    border_left: Any = None,
    border_right: Any = None,
    border_top: Any = None,
    border_bottom: Any = None,
) -> Border:
    """Create a Border.

    cell_border=True means thin light-green border on all sides.
    cell_border=False/None means no border, unless an individual side is set.
    A dict may contain left/right/top/bottom/all entries.
    """
    base_spec: Any
    if cell_border is True:
        base_spec = ("thin", DEFAULT_BORDER_COLOR)
    elif isinstance(cell_border, Mapping):
        all_spec = cell_border.get("all", cell_border.get("default", None))
        return Border(
            left=make_side(cell_border.get("left", border_left if border_left is not None else all_spec)),
            right=make_side(cell_border.get("right", border_right if border_right is not None else all_spec)),
            top=make_side(cell_border.get("top", border_top if border_top is not None else all_spec)),
            bottom=make_side(cell_border.get("bottom", border_bottom if border_bottom is not None else all_spec)),
        )
    elif cell_border in (False, None):
        base_spec = None
    else:
        base_spec = cell_border

    return Border(
        left=make_side(border_left if border_left is not None else base_spec),
        right=make_side(border_right if border_right is not None else base_spec),
        top=make_side(border_top if border_top is not None else base_spec),
        bottom=make_side(border_bottom if border_bottom is not None else base_spec),
    )


def parse_text_style(text_style: Any) -> dict[str, Any]:
    """Parse text_style aliases into font kwargs.

    Examples:
      text_style=("bold", "italic", "red")
      text_style={"bold": True, "italic": True, "color": "red"}
    """
    result: dict[str, Any] = {}
    if text_style is None:
        return result

    if isinstance(text_style, Mapping):
        for key, value in text_style.items():
            k = str(key).lower()
            if k in {"color", "colour", "text_color", "font_color"}:
                result["text_color"] = value
            elif k in {"underline", "underscore"}:
                result["underscore"] = bool(value)
            elif k in {"bold", "italic"}:
                result[k] = bool(value)
            elif k in {"size", "font_size"}:
                result["font_size"] = value
            elif k in {"name", "font_name"}:
                result["font_name"] = value
        return result

    if isinstance(text_style, str):
        items: Iterable[Any] = (text_style,)
    else:
        items = text_style

    for item in items:
        if item is None:
            continue
        if isinstance(item, bool):
            # A bare True is interpreted as bold; a bare False is ignored.
            if item:
                result["bold"] = True
            continue
        token = str(item).strip().lower()
        if token in {"bold", "b"}:
            result["bold"] = True
        elif token in {"italic", "italics", "i"}:
            result["italic"] = True
        elif token in {"underline", "underlined", "underscore", "u"}:
            result["underscore"] = True
        else:
            # Anything else is assumed to be a color name/hex. The color helper
            # will validate it later.
            result["text_color"] = item
    return result


def make_font(
    *,
    bold: bool = False,
    italic: bool = False,
    underscore: bool = False,
    text_color: ColorLike = "black",
    font_name: str | None = None,
    font_size: int | float | None = None,
) -> Font:
    kwargs: dict[str, Any] = {
        "bold": bold,
        "italic": italic,
        "underline": "single" if underscore else None,
        "color": to_argb(text_color, none_ok=False),
    }
    if font_name is not None:
        kwargs["name"] = font_name
    if font_size is not None:
        kwargs["size"] = font_size
    return Font(**kwargs)


def make_alignment(
    alignment: str | None = "right",
    *,
    vertical: str | None = "center",
    wrap_text: bool | None = None,
) -> Alignment:
    if alignment is None:
        horizontal = None
    else:
        key = alignment.strip().lower()
        aliases = {
            "left": "left",
            "center": "center",
            "centre": "center",
            "right": "right",
            "justify": "justify",
            "justified": "justify",
        }
        if key not in aliases:
            raise ValueError("alignment must be one of: left, center, right, justify")
        horizontal = aliases[key]
    return Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap_text)
